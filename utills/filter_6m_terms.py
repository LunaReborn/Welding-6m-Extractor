#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
清洗 Welding-6m-Extractor 导出的 6M 术语 Excel。

主要过滤：
1) “人”类别里的真实人名：如 作者/审者/主编/姓名上下文中的中文姓名，或 2~4 字疑似中文姓名。
2) “料”类别里的独立化学式：如 CO、NO、CO2、O2、H2O、Al2O3 等。
3) 其他常见噪声：章节号、图表号、纯数字符号、过短英文/单字母区间、JSON/图片残留等。

用法：
    python filter_6m_terms.py 1_6m_terms.xlsx -o 1_6m_terms.cleaned.xlsx --removed 1_6m_terms.removed.xlsx

依赖：
    pip install openpyxl
"""

from __future__ import annotations

import argparse
import copy
import re
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# === 可以按你的语料继续扩展的规则区 ===
ALLOWED_CATEGORIES = {"人", "机", "料", "法", "环", "测"}

# “人”类别中应保留的角色/岗位/群体词。命中这些词时，不按姓名过滤。
PERSON_ROLE_KEYWORDS = {
    "工", "员", "师", "长", "者", "人员", "操作者", "操作工", "焊工", "焊接工", "电焊工",
    "检验员", "检测员", "试验员", "质检员", "技术员", "工程师", "工艺师", "监督", "监理",
    "负责人", "管理人员", "作业人员", "施工人员", "维修人员", "操作人员", "检查人员",
    "无损检测人员", "焊接操作人员", "焊接技术人员", "焊接检验人员",
}

# 中文常见姓氏。用于“2~4 个汉字 + 无岗位含义”的疑似姓名判断。
CHINESE_SURNAMES = set(
    "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜"
    "戚谢邹喻柏水窦章云苏潘葛奚范彭郎鲁韦昌马苗凤花方俞任袁柳酆鲍史唐"
    "费廉岑薛雷贺倪汤滕殷罗毕郝邬安常乐于时傅皮卞齐康伍余元卜顾孟平黄"
    "和穆萧尹姚邵湛汪祁毛禹狄米贝明臧计伏成戴谈宋庞熊纪舒屈项祝董梁杜"
    "阮蓝闵席季麻强贾路娄危江童颜郭梅盛林刁钟徐邱骆高夏蔡田胡凌霍虞万"
    "支柯昝管卢莫经房裘缪干解应宗丁宣邓郁单杭洪包诸左石崔吉龚程嵇邢滑"
    "裴陆荣翁荀羊於惠甄曲家封芮羿储靳汲邴糜松井段富巫乌焦巴弓牧隗山谷"
    "车侯宓蓬全郗班仰秋仲伊宫宁仇栾暴甘斜厉戎祖武符刘景詹束龙叶幸司韶"
    "郜黎蓟薄印宿白怀蒲邰从鄂索咸籍赖卓蔺屠蒙池乔阴鬱胥能苍双闻莘党翟"
    "谭贡劳逄姬申扶堵冉宰郦雍郤璩桑桂濮牛寿通边扈燕冀郏浦尚农温别庄晏"
    "柴瞿阎充慕连茹习宦艾鱼容向古易慎戈廖庾终暨居衡步都耿满弘匡国文寇"
    "广禄阙东殴殳沃利蔚越夔隆师巩厍聂晁勾敖融冷訾辛阚那简饶空曾毋沙乜"
    "养鞠须丰巢关蒯相查后荆红游竺权逯盖益桓公" 
)
COMPOUND_SURNAMES = {
    "欧阳", "太史", "端木", "上官", "司马", "东方", "独孤", "南宫", "万俟", "闻人",
    "夏侯", "诸葛", "尉迟", "公羊", "赫连", "澹台", "皇甫", "宗政", "濮阳", "公冶",
    "太叔", "申屠", "公孙", "慕容", "仲孙", "钟离", "长孙", "宇文", "司徒", "鲜于",
    "司空", "闾丘", "子车", "亓官", "司寇", "巫马", "公西", "颛孙", "壤驷", "公良",
    "漆雕", "乐正", "宰父", "谷梁", "拓跋", "夹谷", "轩辕", "令狐", "段干", "百里",
    "呼延", "东郭", "南门", "羊舌", "微生", "公户", "公玉", "公仪", "梁丘", "公仲",
    "公上", "公门", "公山", "公坚", "左丘", "公伯", "西门", "公祖", "第五", "公乘",
    "贯丘", "公皙", "南荣", "东里", "东宫", "仲长", "子书", "子桑", "即墨", "达奚",
    "褚师", "吴铭",
}

# 人名上下文：术语出现在这些词附近时，通常是书籍/论文元数据，而不是 6M 的“人”。
NAME_CONTEXT_RE = re.compile(
    r"(作者|审者|编者|主编|副主编|编著|著者|责任编辑|审稿|审核|校对|译者|姓名|签名|编写|参编|专家|教授|博士|硕士|院士)"
)

# 焊接领域中更有价值的材料实体词尾/词素。命中时，即使含字母数字也尽量保留。
MATERIAL_KEEP_KEYWORDS = {
    "钢", "铁", "铜", "铝", "镍", "钛", "锌", "镁", "合金", "焊丝", "焊条", "焊剂", "焊料", "药皮",
    "气体", "保护气", "氩气", "氦气", "氧气", "氮气", "二氧化碳", "乙炔", "丙烷", "母材", "填充金属",
    "熔敷金属", "电极材料", "钨极", "钍钨", "铈钨", "镧钨", "不锈钢", "低碳钢", "高强钢",
}

GENERAL_STOP_TERMS = {
    "进行", "要求", "符合", "采用", "确保", "可以", "需要", "相关", "过程", "情况", "使用", "工作",
    "内容", "部分", "方面", "条件", "原因", "问题", "方法", "方式", "影响", "作用", "结果", "规定",
    "图", "表", "公式", "章节", "参考文献", "目录", "附录", "前言",
}

# 元素符号集合，用于识别“纯化学式”。
ELEMENTS = {
    "H", "He", "Li", "Be", "B", "C", "N", "O", "F", "Ne", "Na", "Mg", "Al", "Si", "P", "S", "Cl", "Ar",
    "K", "Ca", "Sc", "Ti", "V", "Cr", "Mn", "Fe", "Co", "Ni", "Cu", "Zn", "Ga", "Ge", "As", "Se", "Br", "Kr",
    "Rb", "Sr", "Y", "Zr", "Nb", "Mo", "Tc", "Ru", "Rh", "Pd", "Ag", "Cd", "In", "Sn", "Sb", "Te", "I", "Xe",
    "Cs", "Ba", "La", "Ce", "Pr", "Nd", "Pm", "Sm", "Eu", "Gd", "Tb", "Dy", "Ho", "Er", "Tm", "Yb", "Lu",
    "Hf", "Ta", "W", "Re", "Os", "Ir", "Pt", "Au", "Hg", "Tl", "Pb", "Bi", "Po", "At", "Rn",
}

SUBSCRIPT_MAP = str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789")


def clean_term(term: Any) -> str:
    term = "" if term is None else str(term)
    term = term.translate(SUBSCRIPT_MAP)
    term = term.strip(" ：:，,。；;、\n\t\r'\"`“”‘’（）()[]【】{}")
    term = re.sub(r"\s+", "", term)
    # 常见 LLM/Markdown 残留
    term = term.replace("$", "")
    return term


def contains_any(text: str, words: set[str]) -> bool:
    return any(w in text for w in words)


def is_all_cjk(text: str) -> bool:
    return bool(re.fullmatch(r"[\u4e00-\u9fff]+", text))


def has_chinese_surname(text: str) -> bool:
    if not text:
        return False
    return text[:2] in COMPOUND_SURNAMES or text[0] in CHINESE_SURNAMES


def looks_like_chinese_real_name(term: str, source_text: str = "") -> bool:
    """偏保守：只在“人”类别中使用。角色/岗位词优先保留。"""
    if contains_any(term, PERSON_ROLE_KEYWORDS):
        return False
    if not is_all_cjk(term):
        return False

    # “作者耿正魏继昆 审者陈树君”这类上下文优先删除。
    if 2 <= len(term) <= 4 and has_chinese_surname(term) and NAME_CONTEXT_RE.search(source_text or ""):
        return True

    # 无上下文时，2~3 字常见姓氏且没有岗位含义，基本可视为人名噪声。
    if 2 <= len(term) <= 3 and has_chinese_surname(term):
        return True

    # 复姓姓名常为 3~4 字。
    if 3 <= len(term) <= 4 and term[:2] in COMPOUND_SURNAMES:
        return True

    return False


def looks_like_english_real_name(term: str, source_text: str = "") -> bool:
    if re.fullmatch(r"[A-Z][a-z]+([·\-\s][A-Z][a-z]+){1,3}", term):
        return True
    if NAME_CONTEXT_RE.search(source_text or "") and re.fullmatch(r"[A-Za-z][A-Za-z·\-\s.]{2,40}", term):
        return True
    return False


def tokenize_formula(formula: str) -> list[str] | None:
    """如果 formula 是纯化学式，返回元素 token；否则返回 None。"""
    s = formula.translate(SUBSCRIPT_MAP)
    s = s.replace("（", "(").replace("）", ")")
    s = re.sub(r"[\s·•]", "", s)
    # 只允许元素符号、数字和括号；像 Q235、A-B、MIG 不会通过。
    if not re.fullmatch(r"[A-Z][A-Za-z0-9()]*", s):
        return None
    tokens = re.findall(r"[A-Z][a-z]?", s)
    if not tokens:
        return None
    if all(t in ELEMENTS for t in tokens) and "".join(re.findall(r"[A-Z][a-z]?|\d+|[()]", s)) == s:
        return tokens
    return None


def looks_like_standalone_formula(term: str) -> bool:
    """识别 CO、NO、CO2、O2、H2O、Al2O3 等独立化学式。"""
    t = term.translate(SUBSCRIPT_MAP)
    t = t.strip()
    if contains_any(t, MATERIAL_KEEP_KEYWORDS):
        return False
    tokens = tokenize_formula(t)
    if not tokens:
        return False
    # 单个元素如 Fe、Ni、Cr 也通常不是完整术语；但避免误杀较长牌号/合金名。
    return len(t) <= 10


def looks_like_noise(term: str) -> bool:
    if len(term) < 2 or len(term) > 40:
        return True
    if term in GENERAL_STOP_TERMS:
        return True
    if re.fullmatch(r"[\d\W_]+", term):
        return True
    if re.fullmatch(r"第[一二三四五六七八九十百千\d]+[章节篇卷]?", term):
        return True
    if re.fullmatch(r"[图表]\d+([\-—–]\d+)?", term, flags=re.I):
        return True
    if re.fullmatch(r"[A-Za-z]([\-—–][A-Za-z])?区?", term):
        return True
    if re.fullmatch(r"[A-Za-z]{1,2}", term):
        return True
    if any(x in term.lower() for x in ("content", "image", ".jpg", ".png", ".pdf", "json")):
        return True
    return False


def filter_reason(row: dict[str, Any]) -> str | None:
    term = clean_term(row.get("术语") or row.get("term"))
    category = clean_term(row.get("类别") or row.get("category"))
    normalized = clean_term(row.get("归一化术语") or row.get("normalized_term"))
    source_text = str(row.get("原文片段") or row.get("source_text") or "")

    if not term:
        return "空术语"
    if category not in ALLOWED_CATEGORIES:
        return f"非法类别：{category}"
    if looks_like_noise(term):
        return "通用噪声/章节图表/过短过长"

    if category == "人":
        if looks_like_chinese_real_name(term, source_text):
            return "疑似真实中文人名"
        if normalized and normalized != term and looks_like_chinese_real_name(normalized, source_text):
            return "归一化术语疑似真实中文人名"
        if looks_like_english_real_name(term, source_text):
            return "疑似真实英文人名"

    if category == "料":
        if looks_like_standalone_formula(term):
            return "料类别中的独立化学式"
        if normalized and normalized != term and looks_like_standalone_formula(normalized):
            return "归一化术语为独立化学式"

    return None


def copy_row_style(src_row, dst_row) -> None:
    for src_cell, dst_cell in zip(src_row, dst_row):
        if src_cell.has_style:
            dst_cell._style = copy.copy(src_cell._style)
        if src_cell.number_format:
            dst_cell.number_format = src_cell.number_format
        if src_cell.alignment:
            dst_cell.alignment = copy.copy(src_cell.alignment)


def rows_to_dicts(ws):
    headers = [clean_term(c.value) for c in ws[1]]
    for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        values = [cell.value for cell in row]
        yield idx, row, dict(zip(headers, values))


def normalize_key(row_dict: dict[str, Any], dedupe_by_doc_chunk: bool = False) -> tuple:
    term = clean_term(row_dict.get("归一化术语") or row_dict.get("术语"))
    category = clean_term(row_dict.get("类别"))
    if dedupe_by_doc_chunk:
        return (term, category, row_dict.get("来源文档"), str(row_dict.get("文本段")))
    return (term, category)


def append_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    fill = PatternFill("solid", fgColor="D9EAF7")
    font = Font(bold=True, color="1F1F1F")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border


def autosize(ws) -> None:
    for col_idx, column in enumerate(ws.columns, start=1):
        max_len = 8
        for cell in column[:300]:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(val), 60))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 8), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def clean_workbook(input_path: Path, output_path: Path, removed_path: Path | None, dedupe_by_doc_chunk: bool) -> None:
    wb_in = load_workbook(input_path, read_only=True, data_only=True)
    ws_in = wb_in.active
    headers = [clean_term(c.value) for c in ws_in[1]]
    if "术语" not in headers or "类别" not in headers:
        raise ValueError(f"未找到必要表头：术语、类别。当前表头为：{headers}")

    wb_keep = Workbook(write_only=True)
    ws_keep = wb_keep.create_sheet(ws_in.title or "6M术语")
    ws_keep.append(headers)

    wb_drop = Workbook(write_only=True)
    ws_drop = wb_drop.create_sheet("removed")
    ws_drop.append(headers + ["过滤原因"])

    seen: OrderedDict[tuple, int] = OrderedDict()
    kept_count = dropped_count = dup_count = 0

    for _, src_row, row_dict in rows_to_dicts(ws_in):
        reason = filter_reason(row_dict)
        if reason is None:
            key = normalize_key(row_dict, dedupe_by_doc_chunk=dedupe_by_doc_chunk)
            if key in seen:
                reason = "重复术语"
                dup_count += 1

        if reason is None:
            values = [cell.value for cell in src_row]
            ws_keep.append(values)
            seen[normalize_key(row_dict, dedupe_by_doc_chunk=dedupe_by_doc_chunk)] = kept_count
            kept_count += 1
        else:
            values = [cell.value for cell in src_row] + [reason]
            ws_drop.append(values)
            dropped_count += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb_keep.save(output_path)
    if removed_path:
        removed_path.parent.mkdir(parents=True, exist_ok=True)
        wb_drop.save(removed_path)

    print(f"输入：{input_path}")
    print(f"保留：{kept_count} 条 -> {output_path}")
    print(f"删除：{dropped_count} 条" + (f" -> {removed_path}" if removed_path else ""))
    print(f"其中重复术语：{dup_count} 条")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="筛选/清洗 Welding-6m-Extractor 生成的 6M 术语 Excel")
    parser.add_argument("input", help="输入 xlsx，例如 1_6m_terms.xlsx")
    parser.add_argument("-o", "--output", default=None, help="清洗后的 xlsx；默认在原文件名后加 .cleaned.xlsx")
    parser.add_argument("--removed", default=None, help="被过滤记录报告 xlsx；默认在原文件名后加 .removed.xlsx")
    parser.add_argument(
        "--dedupe-by-doc-chunk",
        action="store_true",
        help="去重时保留同一术语在不同来源文档/文本段中的记录；默认只按 归一化术语+类别 去重",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_path = Path(args.input)
    output_path = Path(args.output) if args.output else input_path.with_suffix(".cleaned.xlsx")
    removed_path = Path(args.removed) if args.removed else input_path.with_suffix(".removed.xlsx")
    clean_workbook(input_path, output_path, removed_path, args.dedupe_by_doc_chunk)


if __name__ == "__main__":
    main()
