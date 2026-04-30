import json
from pathlib import Path

FIELDS = ["term", "category", "normalized_term", "source_text", "document", "chunk_id", "confidence", "reason"]
HEADERS = ["术语", "类别", "归一化术语", "原文片段", "来源文档", "文本段", "置信度", "说明"]


def export_json(records: list[dict], path: str | Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")


def export_xlsx(records: list[dict], path: str | Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("导出 Excel 需要安装 openpyxl：pip install openpyxl") from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "6M术语"
    ws.append(HEADERS)

    for record in records:
        ws.append([record.get(field, "") for field in FIELDS])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True, color="1F1F1F")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    widths = [22, 8, 22, 60, 28, 10, 10, 32]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
